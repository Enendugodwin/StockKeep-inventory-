"""Daily/monthly reporting, catch-up and reconciliation tests."""

from __future__ import annotations

from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from stockkeep.services import analytics, catalog, reports, sales, scheduler

DAY = date(2026, 4, 1)


def _seed(db):
    with db.session() as session:
        rice = catalog.create_product(
            session, name="Rice", selling_price="10", cost_price="6", quantity="10", business_date=DAY
        )
        sales.create_sale(
            session,
            items=[{"product_id": rice.id, "quantity": "2"}],
            payment_method="Cash",
            business_date=DAY,
        )


def test_generate_daily_writes_four_files(db, paths):
    _seed(db)
    outcome = reports.generate_daily(db, DAY, paths.reports_dir)
    assert outcome.status == "generated"
    files = reports.daily_paths(paths.reports_dir, DAY)
    for key in ("transactions_csv", "transactions_xlsx", "movements_csv", "summary_xlsx"):
        assert files[key].is_file(), key

    workbook = load_workbook(files["transactions_xlsx"])
    assert workbook.sheetnames == ["Transactions", "Daily Summary", "Stock Movements"]
    rows = list(workbook["Transactions"].iter_rows(values_only=True))
    assert rows[0][0] == "Transaction ID"
    assert any("Rice" in str(cell) for cell in rows[1])

    # Second run is a no-op unless forced.
    again = reports.generate_daily(db, DAY, paths.reports_dir)
    assert again.status == "skipped"
    forced = reports.generate_daily(db, DAY, paths.reports_dir, force=True)
    assert forced.status == "generated"


def test_catch_up_generates_each_due_day(db, paths):
    _seed(db)
    config = scheduler.ReportConfig(report_root=paths.reports_dir)
    now = datetime(2026, 4, 3, 1, 0)  # 02:00 Lagos on 3 Apr -> 1 & 2 Apr are closed
    summary = scheduler.generate_missing(db, config, now=now)
    assert summary.failed == 0
    assert summary.generated == 2
    assert reports.daily_files_exist(paths.reports_dir, date(2026, 4, 1))
    assert reports.daily_files_exist(paths.reports_dir, date(2026, 4, 2))

    again = scheduler.generate_missing(db, config, now=now)
    assert again.generated == 0


def test_cli_generate_returns_success(db, paths):
    _seed(db)
    config = scheduler.ReportConfig(report_root=paths.reports_dir)
    code = scheduler.cli_generate(db, config, day=DAY)
    assert code == 0
    assert reports.daily_files_exist(paths.reports_dir, DAY)


def test_daily_summary_reconciles_with_database(db, paths):
    _seed(db)
    reports.generate_daily(db, DAY, paths.reports_dir)
    files = reports.daily_paths(paths.reports_dir, DAY)
    workbook = load_workbook(files["summary_xlsx"])
    summary_values = {row[0]: row[1] for row in workbook["Daily Summary"].iter_rows(values_only=True)}
    assert summary_values["Net sales"] == "₦20.00"
    assert summary_values["Units sold"] == "2"
    assert summary_values["Completed sales count"] == "1"


def test_monthly_summary_matches_transactions_and_export(db, paths, tmp_path):
    _seed(db)
    with db.session() as session:
        summary = analytics.monthly_summary(session, 2026, 4)
        assert summary["sales_count"] == 1
        assert summary["net_minor"] == 2000
        assert summary["gross_minor"] == 2000
        assert summary["opening_milli"] == 0
        assert summary["received_milli"] == 0
        assert summary["sold_milli"] == 2000
        assert summary["closing_milli"] == 8000

        xlsx = reports.export_monthly(session, 2026, 4, tmp_path / "monthly.xlsx")
        csv_dir = reports.export_monthly(session, 2026, 4, tmp_path / "monthly.csv", fmt="csv")
        assert xlsx.is_file() and csv_dir.is_dir()
        assert (csv_dir / "executive_summary.csv").is_file()
        workbook = load_workbook(xlsx)
        assert "Executive Summary" in workbook.sheetnames
        assert "Daily Sales" in workbook.sheetnames


def test_export_report_kinds(db, paths, tmp_path):
    _seed(db)
    with db.session() as session:
        out = reports.export_report(session, "sales", DAY, DAY, tmp_path / "sales.xlsx")
        assert out.is_file()
        workbook = load_workbook(out)
        assert workbook.sheetnames == ["Sales Transactions"]
