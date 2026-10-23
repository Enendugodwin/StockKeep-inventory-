"""CSV/Excel import: preview validation, modes and atomicity."""

from __future__ import annotations

import csv
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from stockkeep.errors import ValidationError
from stockkeep.models import ImportBatch, Product
from stockkeep.services import catalog, imports

DAY = date(2026, 4, 1)
HEADERS = [
    "Stock Name",
    "Price",
    "Amount in Inventory",
    "Product ID",
    "SKU",
    "Category",
    "Cost Price",
    "Reorder Level",
    "Unit",
]


def _write_csv(path: Path, rows: list[list[str]]) -> Path:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        writer.writerows(rows)
    return path


def test_templates_are_written(tmp_path):
    csv_path = imports.write_template_csv(tmp_path / "template.csv")
    xlsx_path = imports.write_template_xlsx(tmp_path / "template.xlsx")
    assert csv_path.is_file() and xlsx_path.is_file()
    header = csv_path.read_text(encoding="utf-8-sig").splitlines()[0]
    assert "Stock Name" in header and "Amount in Inventory" in header
    sheet = load_workbook(xlsx_path).active
    assert sheet.cell(row=1, column=1).value == "Stock Name"


def test_preview_classifies_rows(db, tmp_path):
    with db.session() as session:
        catalog.create_product(
            session, name="Existing Rice", selling_price="10", quantity="5", product_code="PRD-1"
        )
    path = _write_csv(
        tmp_path / "stock.csv",
        [
            ["New Beans", "500", "20", "", "", "Food", "300", "5", "pcs"],
            ["Existing Rice", "12", "9", "PRD-1", "", "", "7", "2", "pcs"],
            ["", "100", "5", "", "", "", "", "", ""],
            ["Bad Price", "abc", "5", "", "", "", "", "", ""],
            ["New Beans", "500", "20", "", "", "Food", "300", "5", "pcs"],
        ],
    )
    with db.session() as session:
        preview = imports.preview_file(path, session)
        assert preview.fatal_error is None
        assert preview.total_rows == 5
        assert preview.new_count == 1
        assert preview.update_count == 1
        assert preview.error_count == 2  # blank name + bad price
        assert preview.duplicate_count == 1

        rejected = imports.write_rejected(preview, tmp_path / "rejected.csv")
        text = rejected.read_text(encoding="utf-8-sig")
        assert "Bad Price" in text and "Duplicate" in text


def test_missing_columns_is_fatal(db, tmp_path):
    path = tmp_path / "bad.csv"
    path.write_text("Name,Colour\nRice,White\n", encoding="utf-8")
    with db.session() as session:
        preview = imports.preview_file(path, session)
    assert preview.fatal_error
    assert preview.missing_columns
    assert preview.valid_count == 0


def test_add_new_skips_existing(db, tmp_path):
    with db.session() as session:
        catalog.create_product(
            session, name="Existing", selling_price="10", quantity="5", product_code="PRD-1"
        )
    path = _write_csv(
        tmp_path / "stock.csv",
        [
            ["Existing", "11", "8", "PRD-1", "", "", "", "", "pcs"],
            ["Brand New", "20", "4", "PRD-2", "", "", "10", "", "pcs"],
        ],
    )
    with db.session() as session:
        preview = imports.preview_file(path, session)
        result = imports.apply_import(session, preview, mode="add_new", business_date=DAY)
        assert (result.added, result.updated, result.skipped) == (1, 0, 1)
        existing = session.query(Product).filter(Product.product_code == "PRD-1").one()
        assert existing.quantity_milli == 5000  # untouched
        assert session.query(ImportBatch).count() == 1


def test_update_existing_sets_quantity_with_movement(db, tmp_path):
    with db.session() as session:
        catalog.create_product(
            session, name="Existing", selling_price="10", quantity="5", product_code="PRD-1"
        )
    path = _write_csv(
        tmp_path / "stock.csv",
        [
            ["Existing", "11", "8", "PRD-1", "", "", "", "", "pcs"],
            ["Ignored New", "20", "4", "", "", "", "", "", "pcs"],
        ],
    )
    with db.session() as session:
        preview = imports.preview_file(path, session)
        result = imports.apply_import(session, preview, mode="update_existing", business_date=DAY)
        assert (result.added, result.updated, result.skipped) == (0, 1, 1)
        product = session.query(Product).filter(Product.product_code == "PRD-1").one()
        assert product.quantity_milli == 8000
        assert product.selling_price_minor == 1100


def test_replace_requires_confirmation_and_blocks_bad_rows(db, tmp_path):
    path = _write_csv(tmp_path / "stock.csv", [["Only Item", "5", "2", "", "", "", "", "", "pcs"]])
    with db.session() as session:
        preview = imports.preview_file(path, session)
        with pytest.raises(ValidationError):
            imports.apply_import(session, preview, mode="replace", business_date=DAY)
        assert session.query(Product).count() == 0

    bad = _write_csv(tmp_path / "bad.csv", [["", "5", "2", "", "", "", "", "", "pcs"]])
    with db.session() as session:
        preview = imports.preview_file(bad, session)
        with pytest.raises(ValidationError):
            imports.apply_import(
                session, preview, mode="replace", confirm_replace=True, business_date=DAY
            )


def test_replace_archives_unlisted_products(db, tmp_path):
    with db.session() as session:
        catalog.create_product(session, name="Old A", selling_price="1", quantity="3")
        catalog.create_product(session, name="Old B", selling_price="1", quantity="0")
    path = _write_csv(tmp_path / "stock.csv", [["Only Item", "5", "2", "", "", "", "", "", "pcs"]])
    with db.session() as session:
        preview = imports.preview_file(path, session)
        imports.apply_import(
            session, preview, mode="replace", confirm_replace=True, business_date=DAY
        )
        active = session.query(Product).filter(Product.is_active.is_(True)).all()
        assert [p.name for p in active] == ["Only Item"]
        old_a = session.query(Product).filter(Product.name == "Old A").one()
        assert old_a.quantity_milli == 0 and not old_a.is_active


def test_import_rolls_back_on_failure(db, tmp_path, monkeypatch):
    path = _write_csv(
        tmp_path / "stock.csv",
        [
            ["Good", "5", "5", "", "", "", "", "", "pcs"],
            ["Good Two", "6", "6", "", "", "", "", "", "pcs"],
        ],
    )
    real_create = imports.catalog.create_product
    calls = {"count": 0}

    def failing_create(*args, **kwargs):
        calls["count"] += 1
        if calls["count"] == 2:
            raise RuntimeError("simulated failure")
        return real_create(*args, **kwargs)

    monkeypatch.setattr(imports.catalog, "create_product", failing_create)
    with db.session() as session:
        preview = imports.preview_file(path, session)
        with pytest.raises(RuntimeError):
            imports.apply_import(session, preview, mode="upsert", business_date=DAY)
    with db.session() as session:
        assert session.query(Product).count() == 0
        assert session.query(ImportBatch).count() == 0
