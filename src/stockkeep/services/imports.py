"""CSV/Excel stock import: templates, validation preview and transactional apply."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from ..atomic import savepoint
from .. import security
from ..errors import ConflictError, StockKeepError, ValidationError
from ..exporting import (
    atomic_path,
    parse_import_extension,
    replace_atomically,
    safe_text,
    write_csv,
)
from ..models import ImportBatch, MOVEMENT_IMPORT, Product
from ..money import from_milli, parse_decimal, to_milli
from ..timeutil import today
from . import audit, catalog, stock
from .report_state import mark_daily_dirty

MAX_ROWS = 50_000

TEMPLATE_COLUMNS = [
    "Stock Name",
    "Price",
    "Amount in Inventory",
    "Product ID",
    "SKU",
    "Category",
    "Cost Price",
    "Reorder Level",
    "Unit",
]

SAMPLE_ROWS = [
    ["Coca-Cola 50cl", "350.00", "120", "PRD-1001", "SKU-COKE-50", "Beverages", "250.00", "24", "pcs"],
    ["Indomie Chicken 70g", "200.00", "300", "PRD-1002", "SKU-IND-70", "Snacks", "150.00", "50", "pack"],
    ["Rice (50kg bag)", "65000.00", "12", "PRD-1003", "SKU-RICE-50", "Grains", "58000.00", "5", "bag"],
]

# canonical key -> accepted header spellings (normalised)
ALIASES = {
    "name": {"stock name", "name", "product name", "item name", "stock", "description"},
    "price": {"price", "selling price", "unit price", "sales price"},
    "quantity": {
        "amount in inventory",
        "quantity",
        "stock quantity",
        "qty",
        "quantity on hand",
        "amount",
        "stock",
    },
    "code": {"product id", "product code", "id", "code", "sku id"},
    "sku": {"sku", "barcode", "sku/barcode", "sku or barcode", "bar code"},
    "category": {"category", "product category", "group"},
    "cost": {"cost price", "cost", "buying price", "purchase price"},
    "reorder": {"reorder level", "reorder", "minimum stock", "min stock", "reorder point"},
    "unit": {"unit", "uom", "unit of measure"},
}
REQUIRED = ("name", "price", "quantity")

CODE_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,39}$")


def _normalise_header(value) -> str:
    text = str(value or "").strip().lower().replace("_", " ").replace("-", " ")
    text = re.sub(r"\s+", " ", text)
    return text.rstrip(":").strip()


def _canonical_map(headers: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, header in enumerate(headers):
        normalised = _normalise_header(header)
        for key, aliases in ALIASES.items():
            if normalised in aliases and key not in mapping:
                mapping[key] = index
    return mapping


@dataclass
class ParsedRow:
    row_number: int
    raw: dict[str, str]
    values: dict[str, str]
    status: str = "new"  # new | update | error | duplicate
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    existing_product_id: int | None = None
    existing_name: str | None = None

    @property
    def is_valid(self) -> bool:
        return self.status in {"new", "update"}

    @property
    def display_name(self) -> str:
        return self.values.get("name", "") or "(blank)"

    def error_text(self) -> str:
        return "; ".join(self.errors)


@dataclass
class ImportPreview:
    path: Path
    headers: list[str]
    rows: list[ParsedRow]
    fatal_error: str | None = None
    missing_columns: list[str] = field(default_factory=list)

    @property
    def total_rows(self) -> int:
        return len(self.rows)

    @property
    def new_count(self) -> int:
        return sum(1 for r in self.rows if r.status == "new")

    @property
    def update_count(self) -> int:
        return sum(1 for r in self.rows if r.status == "update")

    @property
    def error_count(self) -> int:
        return sum(1 for r in self.rows if r.status == "error")

    @property
    def duplicate_count(self) -> int:
        return sum(1 for r in self.rows if r.status == "duplicate")

    @property
    def rejected_count(self) -> int:
        return self.error_count + self.duplicate_count

    @property
    def valid_count(self) -> int:
        return self.new_count + self.update_count

    @property
    def rejected_rows(self) -> list[ParsedRow]:
        return [r for r in self.rows if r.status in {"error", "duplicate"}]

    def counts(self) -> dict[str, int]:
        return {
            "total": self.total_rows,
            "valid": self.valid_count,
            "new": self.new_count,
            "update": self.update_count,
            "invalid": self.error_count,
            "duplicates": self.duplicate_count,
            "rejected": self.rejected_count,
        }


@dataclass
class ImportResult:
    batch_id: int
    added: int
    updated: int
    skipped: int
    rejected: int
    total: int

    def summary(self) -> str:
        return (
            f"Added {self.added}, updated {self.updated}, skipped {self.skipped}, "
            f"rejected {self.rejected}."
        )


def _read_csv(path: Path) -> list[list]:
    raw = path.read_bytes()
    if not raw.strip():
        raise ValidationError("The selected CSV file is empty.", "file")
    text = None
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    if text is None:
        raise ValidationError("The CSV file encoding could not be read.", "file")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = []
    for row in reader:
        if len(rows) > MAX_ROWS:
            raise ValidationError(
                f"The file has more than {MAX_ROWS:,} rows. Split it into smaller files.", "file"
            )
        rows.append([str(cell).strip() if cell is not None else "" for cell in row])
    while rows and not any(str(cell).strip() for cell in rows[-1]):
        rows.pop()
    return rows


def _read_xlsx(path: Path) -> list[list]:
    try:
        workbook = load_workbook(filename=path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError(f"The Excel file could not be opened: {exc}", "file") from None
    try:
        worksheet = workbook.worksheets[0] if workbook.worksheets else None
        if worksheet is None:
            raise ValidationError("The workbook contains no worksheets.", "file")
        rows: list[list] = []
        for values in worksheet.iter_rows(values_only=True):
            if len(rows) > MAX_ROWS:
                raise ValidationError(
                    f"The file has more than {MAX_ROWS:,} rows. Split it into smaller files.", "file"
                )
            row = ["" if value is None else str(value).strip() for value in values]
            rows.append(row)
        while rows and not any(cell.strip() for cell in rows[-1]):
            rows.pop()
        return rows
    finally:
        workbook.close()


def read_table(path: Path | str) -> list[list]:
    path = Path(path)
    suffix = parse_import_extension(path)
    if not path.is_file():
        raise ValidationError("The selected file does not exist.", "file")
    if path.stat().st_size == 0:
        raise ValidationError("The selected file is empty.", "file")
    return _read_csv(path) if suffix == ".csv" else _read_xlsx(path)


def _match_existing(
    session: Session, *, code: str | None, sku: str | None, name: str
) -> tuple[Product | None, str | None]:
    if code:
        match = (
            session.query(Product)
            .filter(func.lower(Product.product_code) == code.lower())
            .one_or_none()
        )
        if match:
            if sku:
                sku_match = (
                    session.query(Product).filter(func.lower(Product.sku) == sku.lower()).one_or_none()
                )
                if sku_match and sku_match.id != match.id:
                    return None, (
                        f"Product ID {code} and SKU {sku} belong to different existing products."
                    )
            return match, None
    if sku:
        match = (
            session.query(Product).filter(func.lower(Product.sku) == sku.lower()).one_or_none()
        )
        if match:
            return match, None
    if not code and not sku:
        matches = (
            session.query(Product)
            .filter(func.lower(Product.name) == name.lower())
            .order_by(Product.id.asc())
            .all()
        )
        if len(matches) == 1:
            return matches[0], None
        if len(matches) > 1:
            return None, "Multiple existing products share this name; provide a Product ID or SKU."
    return None, None


def preview_file(path: Path | str, session: Session) -> ImportPreview:
    path = Path(path)
    preview = ImportPreview(path=path, headers=[], rows=[])
    try:
        table = read_table(path)
    except ValidationError as exc:
        preview.fatal_error = exc.message
        return preview
    except Exception as exc:  # malformed file
        preview.fatal_error = f"The file could not be parsed: {exc}"
        return preview

    if not table:
        preview.fatal_error = "The file contains no data."
        return preview
    headers = [str(h).strip() for h in table[0]]
    preview.headers = headers
    mapping = _canonical_map(headers)
    missing = [label for key, label in (("name", "Stock Name"), ("price", "Price"), ("quantity", "Amount in Inventory")) if key not in mapping]
    if missing:
        preview.missing_columns = missing
        preview.fatal_error = "Missing required column(s): " + ", ".join(missing)
        return preview

    seen_identities: dict[str, int] = {}
    seen_product_ids: dict[int, int] = {}
    for index, raw_values in enumerate(table[1:], start=2):
        raw = {}
        for column_index, header in enumerate(headers):
            value = raw_values[column_index] if column_index < len(raw_values) else ""
            raw[header or f"Column {column_index + 1}"] = value
        if not any(str(value).strip() for value in raw_values):
            continue
        values = {
            key: (raw_values[col] if col < len(raw_values) else "")
            for key, col in mapping.items()
        }
        for key in ALIASES:
            values.setdefault(key, "")
        row = ParsedRow(row_number=index, raw=raw, values=values)
        preview.rows.append(row)
        _validate_row(session, row, seen_identities, seen_product_ids)
    if not preview.rows:
        preview.fatal_error = "The file has headers but no data rows."
    return preview


def _validate_row(session: Session, row: ParsedRow, seen: dict[str, int], seen_ids: dict[int, int]) -> None:
    values = row.values
    name = str(values.get("name", "")).strip()
    values["name"] = name
    if not name:
        row.errors.append("Stock Name is required.")
    elif len(name) > 240:
        row.errors.append("Stock Name must be 240 characters or fewer.")

    price = None
    try:
        price = parse_decimal(values.get("price"), field="Price")
        if price < 0:
            row.errors.append("Price cannot be negative.")
            price = None
        elif price == 0:
            row.warnings.append("Price is zero.")
    except ValidationError as exc:
        row.errors.append(exc.message)

    quantity = None
    try:
        quantity = parse_decimal(values.get("quantity"), field="Amount in Inventory")
        if quantity < 0:
            row.errors.append("Amount in Inventory cannot be negative.")
            quantity = None
    except ValidationError as exc:
        row.errors.append(exc.message)

    cost = None
    raw_cost = values.get("cost", "")
    if str(raw_cost).strip() == "":
        row.warnings.append("Cost Price not supplied; cost valuation will be incomplete.")
        cost = parse_decimal(0)
    else:
        try:
            cost = parse_decimal(raw_cost, field="Cost Price")
            if cost < 0:
                row.errors.append("Cost Price cannot be negative.")
                cost = None
        except ValidationError as exc:
            row.errors.append(exc.message)

    reorder = parse_decimal(0)
    if str(values.get("reorder", "")).strip():
        try:
            reorder = parse_decimal(values.get("reorder"), field="Reorder Level")
            if reorder < 0:
                row.errors.append("Reorder Level cannot be negative.")
        except ValidationError as exc:
            row.errors.append(exc.message)

    code = str(values.get("code", "")).strip() or None
    if code and not CODE_RE.match(code):
        row.errors.append("Product ID may contain letters, numbers, dots, dashes and underscores (max 40).")
    sku = str(values.get("sku", "")).strip() or None
    if sku and len(sku) > 80:
        row.errors.append("SKU / barcode must be 80 characters or fewer.")
    unit = str(values.get("unit", "")).strip()
    if not unit:
        unit = "pcs"
        row.warnings.append("Unit not supplied; defaulted to pcs.")
    if len(unit) > 24:
        row.errors.append("Unit must be 24 characters or fewer.")
    values["unit"] = unit
    values["code"] = code or ""
    values["sku"] = sku or ""
    values["category"] = str(values.get("category", "")).strip()

    # Duplicate detection within the file.
    identities = []
    if code:
        identities.append(f"code:{code.lower()}")
    if sku:
        identities.append(f"sku:{sku.lower()}")
    if not identities:
        identities.append(f"name:{name.lower()}::{unit.lower()}")
    duplicate_of = None
    for identity in identities:
        if identity in seen:
            duplicate_of = seen[identity]
            break
    if duplicate_of is not None:
        row.status = "duplicate"
        row.errors.append(f"Duplicate row in file (matches row {duplicate_of}).")
    for identity in identities:
        seen.setdefault(identity, row.row_number)

    existing, conflict = _match_existing(session, code=code, sku=sku, name=name)
    if conflict:
        row.errors.append(conflict)
    if existing is not None:
        first_row = seen_ids.get(existing.id)
        if first_row is not None and first_row != row.row_number:
            row.status = "duplicate"
            row.errors.append(
                f"Duplicate row in file (row {first_row} already updates the same product)."
            )
        seen_ids.setdefault(existing.id, row.row_number)
        row.existing_product_id = existing.id
        row.existing_name = existing.name

    if row.errors:
        if row.status != "duplicate":
            row.status = "error"
    elif existing is not None:
        row.status = "update"
    elif duplicate_of is None:
        row.status = "new"


def write_template_csv(path: Path | str) -> Path:
    return write_csv(path, TEMPLATE_COLUMNS, [row[:] for row in SAMPLE_ROWS])


def write_template_xlsx(path: Path | str) -> Path:
    target = Path(path)
    if target.suffix.lower() != ".xlsx":
        target = target.with_suffix(".xlsx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Stock"
    sheet.append(TEMPLATE_COLUMNS)
    for row in SAMPLE_ROWS:
        sheet.append(row)
    header_fill = PatternFill("solid", fgColor="0E7C86")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    widths = [28, 12, 20, 14, 18, 16, 12, 15, 10]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(TEMPLATE_COLUMNS))}{len(SAMPLE_ROWS) + 1}"

    instructions = workbook.create_sheet("Instructions")
    lines = [
        ["StockKeep stock import template"],
        [""],
        ["Required columns", "Stock Name, Price, Amount in Inventory"],
        ["Optional columns", "Product ID, SKU, Category, Cost Price, Reorder Level, Unit"],
        ["Import modes", "Add new, update matching, add and update, or replace inventory."],
        ["Notes", "Leave Product ID blank to let StockKeep generate one."],
        ["", "Prices use your configured currency. Amounts may be decimal (e.g. 1.5 kg)."],
        ["", "Existing products are matched by Product ID, then SKU, then exact name."],
        ["", "A row with the same SKU or Product ID twice is rejected as a duplicate."],
    ]
    for line in lines:
        instructions.append(line)
    instructions.column_dimensions["A"].width = 22
    instructions.column_dimensions["B"].width = 86
    instructions["A1"].font = Font(bold=True, size=14)
    instructions["A3"].font = Font(bold=True)
    instructions["A4"].font = Font(bold=True)
    instructions["A5"].font = Font(bold=True)
    instructions["A6"].font = Font(bold=True)

    temp = atomic_path(target)
    try:
        workbook.save(temp)
        replace_atomically(temp, target)
    except Exception:
        temp.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()
    return target


def write_rejected(preview: ImportPreview, path: Path | str) -> Path:
    target = Path(path)
    rows = preview.rejected_rows
    headers = list(preview.headers) + ["Rejected Status", "Errors"]
    data = []
    for row in rows:
        original = [row.raw.get(header, "") for header in preview.headers]
        data.append(original + [row.status, row.error_text()])
    if target.suffix.lower() == ".xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Rejected rows"
        sheet.append(headers)
        for values in data:
            sheet.append([safe_text(value) for value in values])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="B91C1C")
        for index, header in enumerate(headers, start=1):
            width = 48 if header == "Errors" else 20
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
        temp = atomic_path(target)
        try:
            workbook.save(temp)
            replace_atomically(temp, target)
        finally:
            workbook.close()
        return target
    return write_csv(target, headers, data)


@security.requires(security.P_INVENTORY_IMPORT)
@savepoint
def apply_import(
    session: Session,
    preview: ImportPreview,
    *,
    mode: str,
    operator_id: int | None = None,
    business_date: date | None = None,
    confirm_replace: bool = False,
    progress=None,
) -> ImportResult:
    """Apply a validated preview in one database transaction.

    The caller owns the transaction; any exception leaves the database
    untouched.  UI progress callbacks receive (done, total).
    """
    if preview.fatal_error:
        raise ValidationError(preview.fatal_error, "file")
    if mode not in {"add_new", "update_existing", "upsert", "replace"}:
        raise ValidationError("Unknown import mode.", "mode")
    if mode == "replace":
        if not confirm_replace:
            raise ValidationError(
                "Inventory replacement requires explicit confirmation.", "confirm"
            )
        if preview.rejected_count:
            raise ValidationError(
                "Replacement is blocked while the file contains invalid or duplicate rows. "
                "Fix or remove those rows first.",
                "file",
            )
    business_day = business_date or today()
    batch = ImportBatch(
        filename=preview.path.name,
        mode=mode,
        rows_total=preview.total_rows,
        operator_id=operator_id,
        status="running",
    )
    session.add(batch)
    session.flush()

    added = updated = skipped = 0
    seen_product_ids: set[int] = set()
    usable = [row for row in preview.rows if row.is_valid]
    total = len(usable) or 1
    try:
        for index, row in enumerate(usable, start=1):
            existing = (
                session.get(Product, row.existing_product_id)
                if row.existing_product_id
                else None
            )
            if mode == "add_new" and existing is not None:
                skipped += 1
                _progress(progress, index, total)
                continue
            if mode == "update_existing" and existing is None:
                skipped += 1
                _progress(progress, index, total)
                continue

            values = row.values
            price = parse_decimal(values.get("price"))
            quantity = parse_decimal(values.get("quantity"))
            cost = parse_decimal(values.get("cost") or 0)
            reorder = parse_decimal(values.get("reorder") or 0)
            code = values.get("code") or None
            sku = values.get("sku") or None
            category = values.get("category") or None
            unit = values.get("unit") or "pcs"

            if existing is None:
                product = catalog.create_product(
                    session,
                    name=values.get("name"),
                    selling_price=price,
                    cost_price=cost,
                    quantity=quantity,
                    reorder_level=reorder,
                    unit=unit,
                    sku=sku,
                    category=category,
                    product_code=code,
                    operator_id=operator_id,
                    business_date=business_day,
                )
                added += 1
                seen_product_ids.add(product.id)
            else:
                catalog.update_product(
                    session,
                    existing.id,
                    name=values.get("name"),
                    selling_price=price,
                    cost_price=cost,
                    reorder_level=reorder,
                    unit=unit,
                    sku=sku if values.get("sku") else None,
                    category=category,
                    operator_id=operator_id,
                )
                target_milli = to_milli(quantity, field="Amount in Inventory")
                if target_milli != existing.quantity_milli:
                    stock.apply_movement(
                        session,
                        existing,
                        target_milli - existing.quantity_milli,
                        MOVEMENT_IMPORT,
                        unit_cost_minor=existing.cost_price_minor,
                        reference_type="import_batch",
                        reference_id=batch.id,
                        reason="Quantity set by stock import",
                        operator_id=operator_id,
                        business_date=business_day,
                    )
                updated += 1
                seen_product_ids.add(existing.id)
            _progress(progress, index, total)

        if mode == "replace":
            others = session.query(Product).filter(Product.is_active.is_(True)).all()
            for product in others:
                if product.id in seen_product_ids:
                    continue
                if product.quantity_milli:
                    stock.apply_movement(
                        session,
                        product,
                        -product.quantity_milli,
                        MOVEMENT_IMPORT,
                        unit_cost_minor=product.cost_price_minor,
                        reference_type="import_batch",
                        reference_id=batch.id,
                        reason="Inventory replaced by import",
                        operator_id=operator_id,
                        business_date=business_day,
                    )
                product.is_active = False
                skipped += 0
    except Exception as exc:
        batch.status = "failed"
        batch.error = str(exc)
        raise

    rejected = preview.rejected_count
    batch.rows_added = added
    batch.rows_updated = updated
    batch.rows_skipped = skipped
    batch.rows_rejected = rejected
    batch.status = "completed"
    audit.record(
        session,
        "inventory.imported",
        "import_batch",
        batch.id,
        operator_id=operator_id,
        details={
            "filename": preview.path.name,
            "mode": mode,
            "added": added,
            "updated": updated,
            "skipped": skipped,
            "rejected": rejected,
        },
    )
    mark_daily_dirty(session, business_day)
    session.flush()
    return ImportResult(
        batch_id=batch.id,
        added=added,
        updated=updated,
        skipped=skipped,
        rejected=rejected,
        total=preview.total_rows,
    )


def _progress(callback, done: int, total: int) -> None:
    if callback is None:
        return
    try:
        callback(done, total)
    except Exception:
        pass


def recent_imports(session: Session, limit: int = 20) -> list[ImportBatch]:
    return (
        session.query(ImportBatch)
        .order_by(ImportBatch.created_at.desc(), ImportBatch.id.desc())
        .limit(limit)
        .all()
    )
